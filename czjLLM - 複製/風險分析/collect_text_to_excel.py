import os
import argparse
import pandas as pd
import hashlib
from openpyxl import Workbook
from openpyxl.styles import Alignment

def calculate_hash(content):
    """
    計算內容的 MD5 哈希值
    """
    return hashlib.md5(content.encode('utf-8')).hexdigest()

def find_original_content(file_path, source_dir, target_dir):
    """
    根據目標檔案路徑，尋找對應的原始檔案
    """
    try:
        # 計算相對路徑，保留完整的子目錄結構
        relative_path = os.path.relpath(file_path, start=target_dir)
        # 去掉處理後的後綴，如 "_Risk Indicator Extraction"
        base_name = os.path.splitext(relative_path)[0].rsplit('_', 1)[0] + ".txt"
        # 使用原始檔根路徑替換目標檔根路徑
        original_path = os.path.join(source_dir, base_name)

        if os.path.exists(original_path):
            with open(original_path, "r", encoding="utf-8") as f:
                original_content = f.read()
                original_word_count = len(original_content)
                return original_content, original_word_count
        else:
            print(f"⚠️ 無法找到原始檔案：{original_path}")
            return "", 0
    except Exception as e:
        print(f"❌ 找原始檔時出錯：{e}")
        return "", 0

def collect_text_files(target_dir, source_dir=""):
    data = []
    include_original = False

    # 自動判斷處理前的原始檔目錄是否存在
    if source_dir and not os.path.exists(source_dir):
        print(f"⚠️ 指定的「處理前的原始檔目錄」不存在：{source_dir}，將跳過原始內容匹配。")
        source_dir = ""
    elif source_dir:
        include_original = True
        print(f"✅ 已找到「處理前的原始檔目錄」：{source_dir}，將包含原始內容。")

    # 遞迴遍歷資料夾及其子資料夾
    for root, _, files in os.walk(target_dir):
        for file in files:
            if file.endswith(".txt"):
                file_path = os.path.join(root, file)
                main_filename = os.path.splitext(os.path.basename(file))[0]  
                parent_dir = os.path.basename(os.path.dirname(file_path))  
                try:
                    with open(file_path, "r", encoding="utf-8") as f:
                        content = f.read()
                    word_count = len(content)  
                    content_hash = calculate_hash(content)  
                    original_content = ""
                    original_word_count = 0
                    if include_original:
                        original_content, original_word_count = find_original_content(file_path, source_dir, target_dir)
                    data.append((file_path, parent_dir, main_filename, word_count, content_hash, content, original_word_count, original_content))
                    print(f"✔ 成功讀取：{file_path}")
                except Exception as e:
                    print(f"❌ 無法讀取：{file_path}，錯誤：{e}")
    return data

def save_to_excel(data, output_path="text_files.xlsx", include_original=False):
    data.sort(key=lambda x: (-x[3], x[4]))

    wb = Workbook()
    ws = wb.active
    ws.title = "Text Files"
    
    headers = ["File Path", "Parent Directory", "Main Filename", "Word Count", "Content Hash", "Content"]
    if include_original:
        headers.extend(["Original Word Count", "Original Content"])
    ws.append(headers)

    for col in ws.iter_cols(min_row=1, max_row=1):
        for cell in col:
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in data:
        ws.append(row)
        for cell in ws[ws.max_row]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    for col in ws.columns:
        max_length = max(len(str(cell.value)) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(15, min(50, max_length + 2))

    try:
        wb.save(output_path)
        print(f"✅ 成功儲存至 Excel：{output_path}")
    except Exception as e:
        print(f"❌ 無法儲存至 Excel，錯誤：{e}")

def display_usage_tips():
    print("\n⚙️  **使用說明：**")
    print("1. 命令列操作：")
    print("   python script.py --target-dir 「處理後的文本目標目錄」 [--source-dir 「處理前的原始檔目錄」] [--output result.xlsx]")
    print("2. 互動模式操作：")
    print("   直接執行程式，會逐步提示輸入路徑。\n")

def main():
    parser = argparse.ArgumentParser(description="文本收集及匯出至 Excel")
    parser.add_argument("--target-dir", help="「處理後的文本目標目錄」")
    parser.add_argument("--source-dir", default="", help="「處理前的原始檔目錄」")
    parser.add_argument("--output", default="text_files.xlsx", help="輸出 Excel 檔案路徑")

    args, unknown = parser.parse_known_args()

    if not args.target_dir:
        display_usage_tips()
        args.target_dir = input("請輸入「處理後的文本目標目錄」：").strip()
        args.source_dir = input("請輸入「處理前的原始檔目錄」（可選）：").strip()
        args.output = input("請輸入輸出 Excel 檔案名稱（預設：text_files.xlsx）：").strip() or "text_files.xlsx"

    if not os.path.exists(args.target_dir):
        print(f"❌ 錯誤：指定的「處理後的文本目標目錄」不存在：{args.target_dir}")
        display_usage_tips()
        return

    data = collect_text_files(args.target_dir, args.source_dir)
    if data:
        save_to_excel(data, output_path=args.output, include_original=bool(args.source_dir))
    else:
        print("⚠️ 沒有找到任何 .txt 檔案")

if __name__ == "__main__":
    main()
