import os
import shutil
import datetime
import pandas as pd
import sqlite3
import json
from typing import Optional
from openpyxl import load_workbook

def read_dataframe(file_path: str, nrows: int = None, table_name: str = None) -> pd.DataFrame:
    """
    根據副檔名自動判斷檔案格式並讀取資料，支援：
    - .xlsx / .xls
    - .csv / .tsv
    - .json
    - .sqlite / .db（需提供 table_name）
    """
    if not os.path.isfile(file_path):
        raise FileNotFoundError(f"❌ 找不到檔案：{file_path}")

    ext = os.path.splitext(file_path)[1].lower()
    print(f"📄 正在讀取檔案：{file_path}")

    try:
        if ext in [".xlsx", ".xls"]:
            df = pd.read_excel(file_path, nrows=nrows)

        elif ext == ".csv":
            try:
                df = pd.read_csv(file_path, encoding="utf-8", nrows=nrows)
            except UnicodeDecodeError:
                print("⚠️ UTF-8 失敗，嘗試以 big5 讀取")
                df = pd.read_csv(file_path, encoding="big5", nrows=nrows)

        elif ext == ".tsv":
            try:
                df = pd.read_csv(file_path, sep="\t", encoding="utf-8", nrows=nrows)
            except UnicodeDecodeError:
                print("⚠️ UTF-8 失敗，嘗試以 big5 讀取")
                df = pd.read_csv(file_path, sep="\t", encoding="big5", nrows=nrows)

        elif ext == ".json":
            with open(file_path, "r", encoding="utf-8") as f:
                data = json.load(f)
            df = pd.json_normalize(data) if isinstance(data, list) else pd.DataFrame([data])

        elif ext in [".sqlite", ".db"]:
            if not table_name:
                raise ValueError("❗讀取 SQLite 檔案時需提供 table_name 參數")
            conn = sqlite3.connect(file_path)
            df = pd.read_sql_query(f"SELECT * FROM {table_name}", conn)
            conn.close()

        else:
            raise ValueError(f"❌ 不支援的檔案格式：{ext}")

        print(f"✅ 成功讀取 {len(df)} 筆資料，欄位數：{len(df.columns)}")
        return df

    except Exception as e:
        raise RuntimeError(f"🚫 讀取失敗：{file_path}\n錯誤：{e}")

def write_dataframe(df: pd.DataFrame, file_path: str, table_name: str = "data", index: bool = False) -> None:
    ext = os.path.splitext(file_path)[1].lower()
    print(f"💾 正在儲存至：{file_path}")

    try:
        if ext == ".xlsx":
            df.to_excel(file_path, index=index)

        elif ext == ".csv":
            df.to_csv(file_path, index=index, encoding="utf-8-sig")

        elif ext == ".tsv":
            df.to_csv(file_path, index=index, sep="\t", encoding="utf-8-sig")

        elif ext == ".json":
            df.to_json(file_path, orient="records", force_ascii=False, indent=2)

        elif ext in [".sqlite", ".db"]:
            conn = sqlite3.connect(file_path)
            df.to_sql(table_name, conn, if_exists="replace", index=index)
            conn.close()

        else:
            raise ValueError(f"❌ 不支援的輸出格式：{ext}")

        print("✅ 資料成功儲存")

    except Exception as e:
        raise RuntimeError(f"🚫 儲存失敗：{file_path}\n錯誤：{e}")

def _is_xlsx(path: str) -> bool:
    return str(path).lower().endswith(".xlsx")

def _sanitize_for_excel(df: pd.DataFrame) -> pd.DataFrame:
    """
    - 將 pd.NA/NaN 轉為 None（openpyxl 需要）
    - 將 list/dict 轉為 JSON 字串（避免「array truth value is ambiguous」）
    """
    out = df.copy()
    # 先把 object 欄位中的 list/dict 轉字串
    for col in out.columns:
        if str(out[col].dtype) == "object":
            out[col] = out[col].apply(
                lambda v: json.dumps(v, ensure_ascii=False) if isinstance(v, (list, dict)) else v
            )
    # 將所有空值變 None，讓 openpyxl 能寫
    out = out.where(~out.isna(), None)
    return out

def write_dataframe_sheet(
    df: pd.DataFrame,
    file_path: str,
    sheet_name: str = "Sheet1",
    index: bool = False
) -> None:
    """
    只覆蓋 Excel 檔案中的某個工作表，保留其它 sheet 不變。
    若 file_path 不是 .xlsx，退回成整檔覆寫（呼叫既有 write_dataframe）。

    Parameters
    ----------
    df : pd.DataFrame
        要寫入的資料
    file_path : str
        目標 .xlsx 檔路徑
    sheet_name : str
        要覆蓋的工作表名稱
    index : bool
        是否寫入索引
    """
    os.makedirs(os.path.dirname(os.path.abspath(file_path)), exist_ok=True)

    if not _is_xlsx(file_path):
        # 非 xlsx 無「多 sheet」概念，保持舊行為
        write_dataframe(df, file_path, index=index)
        return

    df_to_write = _sanitize_for_excel(df)

    exists = os.path.isfile(file_path)

    # 主要路徑：新 pandas（支援 if_sheet_exists）
    try:
        mode = "a" if exists else "w"
        with pd.ExcelWriter(
            file_path,
            engine="openpyxl",
            mode=mode,
            **({"if_sheet_exists": "replace"} if exists else {})
        ) as w:
            df_to_write.to_excel(w, index=index, sheet_name=sheet_name)
        return
    except TypeError:
        # 舊 pandas 不支援 if_sheet_exists：先用 openpyxl 刪除目標 sheet 再以附加方式寫入
        if exists:
            try:
                wb = load_workbook(file_path)
                if sheet_name in wb.sheetnames:
                    ws = wb[sheet_name]
                    wb.remove(ws)
                    wb.save(file_path)
                    wb.close()
            except Exception:
                # 就算刪失敗，也嘗試用新 sheet 名覆蓋；下一步會直接寫
                pass

        # 重新用 append 模式寫入（此時 sheet 已被刪除或不存在）
        with pd.ExcelWriter(file_path, engine="openpyxl", mode=("a" if exists else "w")) as w:
            df_to_write.to_excel(w, index=index, sheet_name=sheet_name)
            
def backup_excel_file(file_path: str) -> str:
    """
    備份指定 Excel 檔案至相同目錄下的 _backup_excel 資料夾，
    檔名會加上 timestamp 以防覆蓋。
    Returns 備份後的完整檔案路徑
    """
    backup_dir = os.path.join(os.path.dirname(file_path), "_backup_excel")
    if not os.path.exists(backup_dir):
        os.makedirs(backup_dir)

    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    original_name = os.path.basename(file_path)
    backup_filename = original_name.replace(".xlsx", f"_backup_{timestamp}.xlsx")
    backup_path = os.path.join(backup_dir, backup_filename)

    shutil.copy(file_path, backup_path)
    print(f"🗂 已備份原始 Excel 至：{backup_path}")
    return backup_path

def backup_table_file(file_path: str) -> str:
    """
    備份指定 Excel 或 CSV 檔案至相同目錄下的 _backup_excel 資料夾，
    檔名會加上 timestamp 以防覆蓋。
    Returns 備份後的完整檔案路徑
    """
    ext = os.path.splitext(file_path)[1].lower()
    if ext not in [".xlsx", ".csv"]:
        print(f"⚠️ 不支援備份的檔案類型：{ext}")
        return ""

    backup_dir = os.path.join(os.path.dirname(file_path), "_backup_excel")
    os.makedirs(backup_dir, exist_ok=True)

    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    base_name = os.path.splitext(os.path.basename(file_path))[0]
    backup_filename = f"{base_name}_backup_{timestamp}{ext}"
    backup_path = os.path.join(backup_dir, backup_filename)

    shutil.copy2(file_path, backup_path)
    print(f"📦 已備份檔案為：{backup_path}")
    return backup_path
    
