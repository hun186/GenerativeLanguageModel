# tests/test_excel_sink.py
import pandas as pd
from openpyxl import load_workbook
from czjLLM.sql_mode.runtime.sinks.excel_sink import ExcelSink

def test_excel_sink_preserve_other_sheets(tmp_path):
    xlsx = tmp_path / "out.xlsx"

    # 1) 先建立含兩個分頁：Sheet1 + KeepMe（確保檔案存在）
    with pd.ExcelWriter(xlsx, engine="openpyxl") as w:
        pd.DataFrame({"A":[1]}).to_excel(w, index=False, sheet_name="Sheet1")
        pd.DataFrame({"B":[2]}).to_excel(w, index=False, sheet_name="KeepMe")

    # 2) 只覆蓋 Sheet1 的內容，不動其他分頁（mode="a" 需要檔案已存在）
    df_in = pd.DataFrame({"Text":["x"], "OPA":[""], "row_id":[0]})
    with pd.ExcelWriter(xlsx, engine="openpyxl", mode="a", if_sheet_exists="replace") as w:
        df_in.to_excel(w, index=False, sheet_name="Sheet1")

    # 模擬你的 ctx
    class Ctx: pass
    ctx = Ctx()
    ctx.output_col = "OPA"; ctx.input_col = "Text"; ctx.record_update_time = True
    ctx.flatten_json_output = False; ctx.primary_key = "row_id"

    # 3) ExcelSink 寫回同一個檔案（或你也可用 output_path 指向另一個檔案）
    sink = ExcelSink(ctx, input_path=str(xlsx), output_path=str(xlsx),
                     flush_main_every=1, flush_sub_every=1)
    sink.open()
    sink.write_main(pk=0, result_text="ok", update_time="2025-01-01 00:00:00")
    sink.close()

    # 4) 驗證其他分頁存在
    wb = load_workbook(xlsx)
    assert "KeepMe" in wb.sheetnames, "other sheet should be preserved"
